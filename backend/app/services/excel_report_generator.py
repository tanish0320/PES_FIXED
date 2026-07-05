"""
SENTINEL Excel Report Generator - Data-Centric Format
=====================================================
Transforms investigation data into professional analytical Excel workbooks.
Designed for auditors, forensic accountants, and investigators.

Format:
- Rows: Individual transactions, entities, or records
- Columns: Feature names, metrics, amounts, dates
- Proper tabular format for data analysis and pivot tables
"""

import logging
import io
from datetime import datetime
from typing import Dict, Any, List, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ExcelReportGenerator:
    """Generates data-centric Excel workbooks for investigation analysis."""

    HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
    HEADER_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    ALT_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

    @staticmethod
    def format_inr(value: float) -> str:
        """Format number as Indian Rupees."""
        try:
            if value is None:
                return ""
            val_int = int(value)
            s = str(val_int)
            if len(s) <= 3:
                return "₹{}".format(s)
            last_three = s[-3:]
            other_parts = s[:-3]
            groups = []
            while other_parts:
                groups.append(other_parts[-2:])
                other_parts = other_parts[:-2]
            groups.reverse()
            formatted = ",".join(groups) + "," + last_three
            return "₹{}".format(formatted)
        except Exception:
            return "₹{:,.2f}".format(float(value)) if value else ""

    @classmethod
    def generate(cls, case: Dict[str, Any], report: Dict[str, Any], transactions: List[Dict[str, Any]]) -> bytes:
        """
        Generate Excel workbook with data-centric format.

        Args:
            case: Case metadata
            report: Investigation report
            transactions: List of transactions

        Returns:
            Bytes of Excel workbook
        """
        case_id = case.get("account_id", "UNKNOWN")
        logger.info("[EXCEL] Generating data-centric workbook for case: {}".format(case_id))

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Case Summary"

            # Build worksheets
            logger.debug("[EXCEL] Building Case Summary sheet")
            cls._build_case_summary(ws, case, report)

            logger.debug("[EXCEL] Building Transactions sheet")
            cls._build_transactions(wb, transactions)

            logger.debug("[EXCEL] Building Money Flow sheet")
            cls._build_money_flow(wb, case, report, transactions)

            logger.debug("[EXCEL] Building Entities sheet")
            cls._build_entities(wb, report)

            logger.debug("[EXCEL] Building Timeline sheet")
            cls._build_timeline(wb, report)

            logger.debug("[EXCEL] Building Risk Analysis sheet")
            cls._build_risk_analysis(wb, case, report)

            logger.info("[EXCEL] Created {} worksheets".format(len(wb.sheetnames)))

            # Save and validate
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            excel_bytes = output.getvalue()

            logger.info("[EXCEL] Workbook saved: {} bytes".format(len(excel_bytes)))

            # Validate
            cls._validate_workbook(excel_bytes, case_id)

            return excel_bytes

        except Exception as e:
            logger.error("[EXCEL] Failed to generate workbook: {}".format(str(e)))
            raise

    @classmethod
    def _validate_workbook(cls, excel_bytes: bytes, case_id: str) -> None:
        """Validate workbook integrity."""
        try:
            test_io = io.BytesIO(excel_bytes)
            test_wb = load_workbook(test_io)
            if not test_wb.sheetnames:
                raise ValueError("Workbook contains no sheets")
            logger.info("[EXCEL] Workbook validation passed")
        except Exception as e:
            logger.error("[EXCEL] Validation failed: {}".format(str(e)))
            raise ValueError("Workbook validation failed: {}".format(str(e)))

    @classmethod
    def _build_case_summary(cls, ws, case: Dict[str, Any], report: Dict[str, Any]):
        """Build case summary sheet."""
        row = 1

        # Title
        ws.merge_cells("A1:D1")
        title = ws["A1"]
        title.value = "Investigation Case Summary"
        title.font = Font(bold=True, size=14, color="FFFFFF")
        title.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25
        row = 3

        # Case metadata table
        data = [
            ["Case ID", case.get("account_id", "N/A")],
            ["Account Holder", case.get("holder_name", "N/A")],
            ["Risk Score", "{:.1f}%".format(case.get("risk_score", 0))],
            ["Risk Level", case.get("risk_level", "N/A").upper()],
            ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Total Transactions", len(case.get("transactions", []))],
        ]

        for item in data:
            ws.cell(row, 1).value = item[0]
            ws.cell(row, 1).font = Font(bold=True)
            ws.cell(row, 2).value = item[1]
            row += 1

        # Executive summary
        row += 1
        ws.cell(row, 1).value = "Executive Summary"
        ws.cell(row, 1).font = Font(bold=True, size=11)
        row += 1

        summary = report.get("executive_summary", "")
        ws.merge_cells("A{}:D{}".format(row, row + 2))
        ws.cell(row, 1).value = summary
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        row += 4

        # Risk explanation
        ws.cell(row, 1).value = "Risk Explanation"
        ws.cell(row, 1).font = Font(bold=True, size=11)
        row += 1

        risk_exp = report.get("risk_explanation", "")
        ws.merge_cells("A{}:D{}".format(row, row + 2))
        ws.cell(row, 1).value = risk_exp
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 20

    @classmethod
    def _build_transactions(cls, wb: Workbook, transactions: List[Dict[str, Any]]):
        """Build transactions table with proper columns."""
        ws = wb.create_sheet("Transactions")

        # Headers
        headers = [
            "Date", "Time", "Description", "Amount (₹)", "Type",
            "Channel", "Counterparty", "Risk Score (%)"
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col)
            cell.value = header
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = cls.HEADER_ALIGNMENT
            cell.border = cls.HEADER_BORDER

        ws.freeze_panes = "A2"

        # Data rows
        for row_idx, tx in enumerate(transactions, 2):
            ws.cell(row_idx, 1).value = tx.get("date", "")
            ws.cell(row_idx, 2).value = tx.get("time", "")
            ws.cell(row_idx, 3).value = tx.get("description", "")

            amount = tx.get("amount", 0)
            ws.cell(row_idx, 4).value = amount
            ws.cell(row_idx, 4).number_format = '#,##0.00'

            ws.cell(row_idx, 5).value = "Debit" if tx.get("is_debit", True) else "Credit"
            ws.cell(row_idx, 6).value = tx.get("channel", "")
            ws.cell(row_idx, 7).value = tx.get("sender_account" if tx.get("is_debit") else "receiver_account", "")

            risk = tx.get("risk_score", 0)
            ws.cell(row_idx, 8).value = risk
            ws.cell(row_idx, 8).number_format = '0.0'

            # Alternating row colors
            if row_idx % 2 == 0:
                for col in range(1, 9):
                    ws.cell(row_idx, col).fill = cls.ALT_FILL

        # Column widths
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 10
        ws.column_dimensions["F"].width = 12
        ws.column_dimensions["G"].width = 20
        ws.column_dimensions["H"].width = 12

    @classmethod
    def _build_money_flow(cls, wb: Workbook, case: Dict[str, Any], report: Dict[str, Any], transactions: List[Dict[str, Any]]):
        """Build money flow analysis sheet."""
        ws = wb.create_sheet("Money Flow Analysis")

        row = 1

        # Summary table
        headers = ["Metric", "Value (₹)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row, col)
            cell.value = header
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = cls.HEADER_ALIGNMENT
            cell.border = cls.HEADER_BORDER

        row = 2
        flow = report.get("money_flow_summary", {})

        metrics = [
            ["Total Inflow", flow.get("total_inflow", 0)],
            ["Total Outflow", flow.get("total_outflow", 0)],
            ["Net Flow", flow.get("net_flow", 0)],
        ]

        for metric, value in metrics:
            ws.cell(row, 1).value = metric
            ws.cell(row, 1).font = Font(bold=True)
            ws.cell(row, 2).value = value
            ws.cell(row, 2).number_format = '#,##0.00'
            row += 1

        row += 2

        # Beneficiaries table
        ws.cell(row, 1).value = "Top Beneficiaries"
        ws.cell(row, 1).font = Font(bold=True, size=11)
        row += 1

        benef_headers = ["Beneficiary", "Total Received (₹)", "Transaction Count"]
        for col, header in enumerate(benef_headers, 1):
            cell = ws.cell(row, col)
            cell.value = header
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = cls.HEADER_ALIGNMENT
            cell.border = cls.HEADER_BORDER

        row += 1
        for benef in report.get("top_beneficiaries", []):
            ws.cell(row, 1).value = benef.get("name", "")
            ws.cell(row, 2).value = benef.get("total_received", 0)
            ws.cell(row, 2).number_format = '#,##0.00'
            ws.cell(row, 3).value = benef.get("tx_count", 0)
            row += 1

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 20

    @classmethod
    def _build_entities(cls, wb: Workbook, report: Dict[str, Any]):
        """Build entities table."""
        ws = wb.create_sheet("Entities")

        row = 1
        entities = report.get("extracted_entities", {})

        entity_types = [
            ("Names", entities.get("names", [])),
            ("Accounts", entities.get("accounts", [])),
            ("UPI IDs", entities.get("upi_ids", [])),
            ("IFSC Codes", entities.get("ifsc", [])),
            ("Beneficiaries", entities.get("beneficiaries", [])),
            ("Merchants", entities.get("merchants", [])),
            ("Banks", entities.get("banks", [])),
        ]

        for entity_type, entity_list in entity_types:
            if not entity_list:
                continue

            # Section header
            ws.cell(row, 1).value = entity_type
            ws.cell(row, 1).font = Font(bold=True, size=11, color="FFFFFF")
            ws.cell(row, 1).fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            row += 1

            # Entity data
            for entity in entity_list:
                ws.cell(row, 1).value = entity.get("value", "")
                ws.cell(row, 2).value = entity.get("occurrences", 1)
                row += 1

            row += 1

        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 15

    @classmethod
    def _build_timeline(cls, wb: Workbook, report: Dict[str, Any]):
        """Build timeline table."""
        ws = wb.create_sheet("Timeline")

        headers = ["Date", "Time", "Event", "Amount (₹)", "Risk Flag", "Pattern"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col)
            cell.value = header
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = cls.HEADER_ALIGNMENT
            cell.border = cls.HEADER_BORDER

        ws.freeze_panes = "A2"

        row = 2
        for event in report.get("timeline", []):
            ws.cell(row, 1).value = event.get("date", "")
            ws.cell(row, 2).value = event.get("time", "")
            ws.cell(row, 3).value = event.get("event", "")
            ws.cell(row, 4).value = event.get("amount", 0)
            ws.cell(row, 4).number_format = '#,##0.00'
            ws.cell(row, 5).value = "Yes" if event.get("risk_flag") else "No"
            ws.cell(row, 6).value = event.get("pattern", "")

            if row % 2 == 0:
                for col in range(1, 7):
                    ws.cell(row, col).fill = cls.ALT_FILL

            row += 1

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 20

    @classmethod
    def _build_risk_analysis(cls, wb: Workbook, case: Dict[str, Any], report: Dict[str, Any]):
        """Build risk analysis sheet."""
        ws = wb.create_sheet("Risk Analysis")

        row = 1

        # Risk metrics
        headers = ["Metric", "Value"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row, col)
            cell.value = header
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = cls.HEADER_ALIGNMENT
            cell.border = cls.HEADER_BORDER

        row = 2
        metrics = [
            ["Risk Score", "{:.1f}%".format(case.get("risk_score", 0))],
            ["Risk Level", case.get("risk_level", "N/A").upper()],
            ["Detected Patterns", len(report.get("detected_patterns", []))],
            ["High Risk Transactions", len(report.get("high_risk_transactions", []))],
            ["Unique Counterparties", report.get("money_flow_summary", {}).get("unique_counterparties", 0)],
        ]

        for metric, value in metrics:
            ws.cell(row, 1).value = metric
            ws.cell(row, 1).font = Font(bold=True)
            ws.cell(row, 2).value = str(value)
            row += 1

        row += 2

        # Detected patterns
        ws.cell(row, 1).value = "Detected Patterns"
        ws.cell(row, 1).font = Font(bold=True, size=11)
        row += 1

        pattern_headers = ["Pattern Name", "Severity", "Confidence (%)", "Description"]
        for col, header in enumerate(pattern_headers, 1):
            cell = ws.cell(row, col)
            cell.value = header
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = cls.HEADER_ALIGNMENT
            cell.border = cls.HEADER_BORDER

        row += 1
        for pattern in report.get("detected_patterns", []):
            ws.cell(row, 1).value = pattern.get("name", "")
            ws.cell(row, 2).value = pattern.get("severity", "").upper()
            ws.cell(row, 3).value = pattern.get("confidence", 0)
            ws.cell(row, 3).number_format = '0.0'
            ws.cell(row, 4).value = pattern.get("description", "")
            row += 1

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 40
