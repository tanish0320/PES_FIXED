import os
import requests
import openpyxl
import io

API_BASE = "http://localhost:8000"
DATASET_DIR = "c:\\Claude_projects\\PES_NEED_FIXING\\Bank-statements-dataset\\primary"

def test_multi_upload():
    print("--- [TEST] Starting Multi-Statement Upload Verification ---")
    
    # 1. Select sample files
    files_to_upload = [
        "00869354051.pdf",
        "17771917925.pdf",
        "18306700003.pdf"
    ]
    
    upload_payload = []
    opened_files = []
    
    for fname in files_to_upload:
        fpath = os.path.join(DATASET_DIR, fname)
        if not os.path.exists(fpath):
            print(f"[ERROR] Sample file not found: {fpath}")
            return
        
        f = open(fpath, "rb")
        opened_files.append(f)
        upload_payload.append(("files", (fname, f, "application/pdf")))
        
    print(f"Uploading {len(files_to_upload)} files to {API_BASE}/upload ...")
    
    try:
        # Send POST request to /upload
        response = requests.post(f"{API_BASE}/upload", files=upload_payload)
        
        # Close file handles
        for f in opened_files:
            f.close()
            
        if response.status_code != 200:
            print(f"[FAILED] Ingestion failed with status {response.status_code}: {response.text}")
            return
            
        result = response.json()
        case_id = result.get("case_id")
        case = result.get("case", {})
        
        print("\n[SUCCESS] Unified Case Created Successfully!")
        print(f"Case ID: {case_id}")
        print(f"Merged Transactions: {result.get('parser_stats', {}).get('parsed_rows')} rows")
        print(f"Primary Statement Accounts: {case.get('account_ids')}")
        print(f"Source Files status:")
        for fu in case.get("files_uploaded", []):
            print(f" - {fu['filename']}: {fu['status']} ({fu['rows_parsed']} parsed, confidence {fu['confidence']}%)")
            
        # 2. Test Excel Download
        print(f"\nRequesting Excel report for Case {case_id}...")
        report_url = f"{API_BASE}/investigation/{case_id}/report/excel"
        report_response = requests.get(report_url)
        
        if report_response.status_code != 200:
            print(f"[FAILED] Excel export failed with status {report_response.status_code}")
            return
            
        excel_bytes = report_response.content
        print(f"Excel report downloaded: {len(excel_bytes)} bytes")
        
        # 3. Validate worksheets structure
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        print(f"Excel Sheet Names: {sheet_names}")
        
        required_sheets = [
            "Executive Summary",
            "Dashboard",
            "Investigation Summary",
            "Transactions",
            "High Risk Transactions",
            "Detected Patterns",
            "Money Flow",
            "Extracted Entities",
            "Timeline",
            "Beneficiaries",
            "Graph Statistics",
            "Parser Statistics",
            "Recommendations",
            "Uploaded Files",
            "Parser Summary",
            "File Statistics",
            "Cross File Links",
            "Shared Entities"
        ]
        
        missing = [s for s in required_sheets if s not in sheet_names]
        if missing:
            print(f"[FAILED] Missing required worksheets: {missing}")
        else:
            print("[SUCCESS] All 18 worksheets are present and verified in the Excel report!")
            
        # 4. Check "Transactions" worksheet column counts
        ws_txs = wb["Transactions"]
        header_vals = [ws_txs.cell(4, c).value for c in range(1, 12)]
        print(f"Transactions header row 4 columns: {header_vals}")
        if "Source File" in header_vals:
            print("[SUCCESS] 'Source File' column is present in Transactions worksheet!")
        else:
            print("[FAILED] 'Source File' column was not found in Transactions headers.")
            
    except Exception as e:
        print(f"[ERROR] Test execution failed: {str(e)}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    test_multi_upload()
